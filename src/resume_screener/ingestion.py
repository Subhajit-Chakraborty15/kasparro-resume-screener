"""Resume ingestion: discover files and turn each into clean text + links.

One bad file must never break the batch, so every reader raises a single
``ResumeReadError`` that the pipeline records and moves on from.
"""
from __future__ import annotations

import hashlib
import logging
import re
from pathlib import Path

from .models import RawDocument

log = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}
MIN_TEXT_CHARS = 40


class ResumeReadError(Exception):
    """A resume could not be read (corrupt, encrypted, empty, unsupported...)."""


def discover_files(input_dir: Path) -> tuple[list[Path], list[Path]]:
    """Return (supported, unsupported) files under ``input_dir`` (recursive, sorted)."""
    if not input_dir.is_dir():
        raise NotADirectoryError(f"Input directory not found: {input_dir}")
    supported: list[Path] = []
    unsupported: list[Path] = []
    for p in sorted(input_dir.rglob("*")):
        if not p.is_file() or p.name.startswith((".", "~$")):
            continue
        (supported if p.suffix.lower() in SUPPORTED_EXTENSIONS else unsupported).append(p)
    return supported, unsupported


def read_document(path: Path, max_file_mb: int = 15) -> RawDocument:
    """Read one resume into a :class:`RawDocument`. Raises ResumeReadError."""
    try:
        size = path.stat().st_size
        if size == 0:
            raise ResumeReadError("File is empty (0 bytes)")
        if size > max_file_mb * 1024 * 1024:
            raise ResumeReadError(f"File larger than {max_file_mb} MB")
        data = path.read_bytes()
    except ResumeReadError:
        raise
    except OSError as exc:
        raise ResumeReadError(f"Cannot read file: {exc}") from exc

    suffix = path.suffix.lower()
    try:
        if suffix == ".pdf":
            text, links, pages = _read_pdf(path)
        elif suffix == ".docx":
            text, links, pages = _read_docx(path)
        elif suffix == ".txt":
            text, links, pages = data.decode("utf-8", errors="replace"), [], None
        else:
            raise ResumeReadError(f"Unsupported file type '{suffix}'")
    except ResumeReadError:
        raise
    except Exception as exc:  # corrupt PDFs raise a zoo of exception types
        raise ResumeReadError(f"Unreadable {suffix[1:].upper()}: {type(exc).__name__}: {exc}") from exc

    text = normalize_text(text)
    if len(text) < MIN_TEXT_CHARS:
        raise ResumeReadError("No extractable text (scanned/image-only file? OCR is not supported)")

    return RawDocument(
        path=str(path),
        text=text,
        links=sorted(set(links)),
        file_hash=hashlib.sha256(data).hexdigest(),
        page_count=pages,
    )


# ------------------------------------------------------------------- readers
def _read_pdf(path: Path) -> tuple[str, list[str], int]:
    from pypdf import PdfReader

    reader = PdfReader(str(path), strict=False)
    if reader.is_encrypted:
        try:
            if not reader.decrypt(""):
                raise ResumeReadError("PDF is password protected")
        except ResumeReadError:
            raise
        except Exception as exc:
            raise ResumeReadError(f"PDF is encrypted and cannot be opened: {exc}") from exc

    texts: list[str] = []
    links: list[str] = []
    for page in reader.pages:
        try:
            texts.append(page.extract_text() or "")
        except Exception as exc:  # one bad page should not lose the rest
            log.debug("page extraction failed in %s: %s", path.name, exc)
        links.extend(_pdf_page_links(page))
    text = "\n".join(texts)

    if len(text.strip()) < MIN_TEXT_CHARS:  # fallback extractor
        text = _read_pdf_with_pdfplumber(path) or text
    return text, links, len(reader.pages)


def _pdf_page_links(page) -> list[str]:
    """Hyperlink annotations: 'GitHub' is often anchor text with the URL hidden."""
    out: list[str] = []
    try:
        for annot in page.get("/Annots") or []:
            obj = annot.get_object()
            action = obj.get("/A")
            if action is None:
                continue
            uri = action.get_object().get("/URI")
            if uri:
                out.append(str(uri))
    except Exception:
        pass
    return out


def _read_pdf_with_pdfplumber(path: Path) -> str:
    try:
        import pdfplumber

        with pdfplumber.open(str(path)) as pdf:
            return "\n".join((pg.extract_text() or "") for pg in pdf.pages)
    except Exception as exc:
        log.debug("pdfplumber fallback failed for %s: %s", path.name, exc)
        return ""


def _read_docx(path: Path) -> tuple[str, list[str], None]:
    import docx

    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    links = [
        rel.target_ref
        for rel in document.part.rels.values()
        if rel.reltype.endswith("/hyperlink")
    ]
    return "\n".join(parts), links, None


# ------------------------------------------------------------- normalisation
_BULLET_GLYPHS = re.compile(r"[\x7f\uf0b7\uf0a7\uf076\uf0d8\uf0fc\u2022\u25cf\u25aa\u25e6\u2023\u2219\u25a0\u25ba\u27a2\u27a4]")
_CID = re.compile(r"\(cid:\d+\)")


def normalize_text(text: str) -> str:
    """Unify bullets/whitespace so downstream regexes behave consistently."""
    text = text.replace("\x00", " ").replace("\u00a0", " ").replace("\r\n", "\n").replace("\r", "\n")
    text = _CID.sub("•", text)
    text = _BULLET_GLYPHS.sub("•", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return "\n".join(line.strip() for line in text.split("\n")).strip()
